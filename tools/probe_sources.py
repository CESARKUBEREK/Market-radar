#!/usr/bin/env python3
"""Comprueba si una fuente candidata responde con datos (Etapa 0).

Solo lee: no escribe nada en la base de datos. Cada comprobación imprime una
anotación de GitHub (::notice:: o ::warning::) con el resultado, para poder
leerlo desde la API sin descargar los registros.
"""

from __future__ import annotations

import os
import sys
from datetime import date, timedelta

import requests

UA = {"User-Agent": "Mozilla/5.0 (market-radar personal data probe)"}
START = (date.today() - timedelta(days=30)).isoformat()


def report(name: str, ok: bool, detail: str) -> None:
    level = "notice" if ok else "warning"
    print(f"::{level}::{'OK ' if ok else 'NO '} {name}: {detail}"[:900])


def tiingo_daily(ticker: str) -> None:
    key = os.environ.get("TIINGO_API_KEY", "").strip()
    try:
        r = requests.get(f"https://api.tiingo.com/tiingo/daily/{ticker}/prices",
                         headers={"Authorization": f"Token {key}"},
                         params={"startDate": START}, timeout=30)
        data = r.json() if r.headers.get("content-type", "").startswith("application/json") else r.text
        if r.ok and isinstance(data, list) and data:
            report(f"tiingo daily {ticker}", True, f"{len(data)} días, último {data[-1]['date'][:10]} cierre {data[-1].get('close')}")
        else:
            report(f"tiingo daily {ticker}", False, f"HTTP {r.status_code} {str(data)[:200]}")
    except Exception as exc:
        report(f"tiingo daily {ticker}", False, str(exc)[:200])


def tiingo_fx(pair: str) -> None:
    key = os.environ.get("TIINGO_API_KEY", "").strip()
    try:
        r = requests.get(f"https://api.tiingo.com/tiingo/fx/{pair}/prices",
                         headers={"Authorization": f"Token {key}"},
                         params={"startDate": START, "resampleFreq": "1day"}, timeout=30)
        data = r.json() if r.headers.get("content-type", "").startswith("application/json") else r.text
        if r.ok and isinstance(data, list) and data:
            first, last = data[0]["date"][:10], data[-1]
            report(f"tiingo fx {pair}", True, f"{len(data)} días desde {first}, último {last['date'][:10]} cierre {last.get('close')}")
        else:
            report(f"tiingo fx {pair}", False, f"HTTP {r.status_code} {str(data)[:200]}")
    except Exception as exc:
        report(f"tiingo fx {pair}", False, str(exc)[:200])


def tiingo_fx_history(pair: str) -> None:
    """Profundidad del historial diario de un par en Tiingo FX."""
    key = os.environ.get("TIINGO_API_KEY", "").strip()
    r = requests.get(f"https://api.tiingo.com/tiingo/fx/{pair}/prices",
                     headers={"Authorization": f"Token {key}"},
                     params={"startDate": "2020-01-01", "endDate": "2020-03-31", "resampleFreq": "1day"}, timeout=60)
    data = r.json() if r.ok else r.text
    ok = r.ok and isinstance(data, list) and bool(data)
    report(f"tiingo fx historial {pair} 2020T1", ok,
           f"{len(data)} días, primero {data[0]['date'][:10]} cierre {data[0].get('close')}" if ok else f"HTTP {r.status_code} {str(data)[:200]}")


def yahoo(symbol: str) -> None:
    try:
        r = requests.get(f"https://query1.finance.yahoo.com/v8/finance/chart/{symbol}",
                         params={"range": "1mo", "interval": "1d"}, headers=UA, timeout=30)
        if not r.ok:
            report(f"yahoo {symbol}", False, f"HTTP {r.status_code}")
            return
        res = r.json()["chart"]["result"][0]
        closes = [c for c in res["indicators"]["quote"][0]["close"] if c is not None]
        report(f"yahoo {symbol}", bool(closes),
               f"{len(closes)} cierres, último {closes[-1] if closes else None}, "
               f"divisa {res['meta'].get('currency')}, bolsa {res['meta'].get('exchangeName')}")
    except Exception as exc:
        report(f"yahoo {symbol}", False, str(exc)[:200])


def euronext(isin_mic: str) -> None:
    url = f"https://live.euronext.com/en/ajax/AwlHistoricalPrice/getFullDownloadAjax/{isin_mic}"
    params = {"format": "csv", "decimal_separator": ".", "date_form": "Y-m-d", "op": "",
              "adjusted": "", "base100": "", "startdate": START, "enddate": date.today().isoformat()}
    for how in ("get", "post"):
        try:
            if how == "get":
                r = requests.get(url, params=params, headers=UA, timeout=30)
            else:
                r = requests.post(url, params={"format": "csv"}, data=params, headers=UA, timeout=30)
            lines = [l for l in r.text.splitlines() if l.strip()]
            report(f"euronext {how} {isin_mic}", r.ok and len(lines) > 3,
                   f"HTTP {r.status_code}, {len(lines)} líneas; {' | '.join(lines[:5])[:300]}")
        except Exception as exc:
            report(f"euronext {how} {isin_mic}", False, str(exc)[:200])


def imf() -> None:
    """Reservas de oro (IRFCL, onzas troy) y deuda/PIB (WEO)."""
    import csv
    import io
    import re
    gold = ("https://api.imf.org/external/sdmx/2.1/data/IRFCL/"
            "USA+DEU+ITA+FRA+CHN+RUS+IND+CHE+JPN.IRFCLDT1_IRFCL56V_FTO..M")
    try:
        r = requests.get(gold, params={"startPeriod": "2020-01"},
                         headers={**UA, "Accept": "application/vnd.sdmx.data+csv;version=1.0.0"}, timeout=90)
        rows = list(csv.DictReader(io.StringIO(r.text)))
        keys = ["COUNTRY", "INDICATOR", "SECTOR", "FREQUENCY", "TIME_PERIOD", "OBS_VALUE", "SCALE",
                "UNIT", "PUBLICATION_DATE", "UPDATE_DATE", "STATUS"]
        del keys
        by: dict = {}
        for x in rows:
            by.setdefault((x.get("COUNTRY"), x.get("SECTOR")), []).append(x)
        parts = []
        for (c, sec), xs in sorted(by.items(), key=lambda kv: str(kv[0])):
            xs.sort(key=lambda x: x["TIME_PERIOD"])
            parts.append(f"{c}/{sec}:{len(xs)}:{xs[-1]['TIME_PERIOD']}={float(xs[-1]['OBS_VALUE']) / 32150.7466:.0f}t")
        report("imf irfcl resumen", bool(parts), " ".join(parts))
    except Exception as exc:
        report("imf irfcl", False, str(exc)[:200])

    debt = [
        ("datamapper", "https://www.imf.org/external/datamapper/api/v1/GGXWDG_NGDP/USA/JPN",
         {"Accept": "application/json"}),
        ("sdmx3 IMF.RES/WEO", "https://api.imf.org/external/sdmx/3.0/data/dataflow/IMF.RES/WEO/~/USA+JPN.GGXWDG_NGDP.A",
         {"Accept": "text/csv"}),
        ("sdmx21 WEO", "https://api.imf.org/external/sdmx/2.1/data/WEO/USA+JPN.GGXWDG_NGDP.A",
         {"Accept": "application/vnd.sdmx.data+csv;version=1.0.0"}),
    ]
    for name, url, h in debt:
        try:
            r = requests.get(url, headers={**UA, **h}, timeout=60)
            lines = r.text.splitlines()
            report(f"imf deuda {name}", r.ok and len(lines) > 1,
                   f"HTTP {r.status_code} {r.headers.get('content-type', '')[:40]}, {len(lines)} líneas; "
                   + " | ".join(ln[:250] for ln in lines[:2]) + " || " + (lines[-1][:250] if lines else ""))
        except Exception as exc:
            report(f"imf deuda {name}", False, str(exc)[:200])
    key = os.environ.get("FRED_API_KEY", "").strip()
    if key:
        got = []
        for iso2 in ("US", "JP", "FR", "GB", "CN", "IN", "AR", "RU"):
            try:
                r = requests.get("https://api.stlouisfed.org/fred/series/observations",
                                 params={"series_id": f"GGGDTA{iso2}A188N", "api_key": key,
                                         "file_type": "json", "observation_start": "2020-01-01"}, timeout=30)
                obs = r.json().get("observations", []) if r.ok else []
                got.append(f"{iso2}:{obs[-1]['date'][:4]}={obs[-1]['value']}" if obs else f"{iso2}:HTTP{r.status_code}")
            except Exception as exc:
                got.append(f"{iso2}:{str(exc)[:40]}")
        report("imf deuda vía FRED (GGGDTA..A188N)", any("=" in g for g in got), " ".join(got))
    try:
        r = requests.get("https://api.imf.org/external/sdmx/2.1/dataflow/all/all/latest",
                         headers={**UA, "Accept": "application/xml"}, timeout=60)
        tags = re.findall(r"<(?:str:|structure:)?Dataflow\b[^>]*>", r.text)
        ids = []
        for t in tags:
            i = re.search(r'\bid="([^"]+)"', t)
            a = re.search(r'agencyID="([^"]+)"', t)
            if i:
                ids.append(f"{a.group(1) if a else '?'}/{i.group(1)}")
        hits = [x for x in ids if any(k in x.upper() for k in ("WEO", "GDD", "DEBT", "GFS", "FM", "HPDD"))]
        report("imf dataflows", bool(ids), f"HTTP {r.status_code}, {len(ids)} flujos; relevantes: {hits[:40]}")
    except Exception as exc:
        report("imf dataflows", False, str(exc)[:200])


def events() -> None:
    """Calendario (fechas de publicación de FRED, reuniones del BCE) y fuentes RSS de noticias."""
    import re
    key = os.environ.get("FRED_API_KEY", "").strip()
    try:
        r = requests.get("https://api.stlouisfed.org/fred/releases/dates", params={
            "api_key": key, "file_type": "json", "realtime_start": date.today().isoformat(),
            "realtime_end": (date.today() + timedelta(days=70)).isoformat(),
            "include_release_dates_with_no_data": "true", "limit": 1000, "sort_order": "asc"}, timeout=60)
        ds = r.json().get("release_dates", [])
        wanted = re.compile(r"Consumer Price|Employment Situation|FOMC|Gross Domestic|Personal Income|"
                            r"Producer Price|Retail|H\.4\.1|Job Openings|Unemployment Insurance", re.I)
        seen: dict = {}
        for d in ds:
            if wanted.search(d.get("release_name", "")):
                seen.setdefault((d["release_id"], d["release_name"]), []).append(d["date"])
        report("fred release dates", bool(seen), " | ".join(f"{i} {n[:35]}: {v[:3]}" for (i, n), v in sorted(seen.items())))
    except Exception as exc:
        report("fred release dates", False, str(exc)[:200])
    try:
        r = requests.get("https://www.ecb.europa.eu/press/calendars/mgcgc/html/index.en.html", headers=UA, timeout=60)
        txt = re.sub(r"<[^>]+>", " ", r.text)
        hits = re.findall(r"(\d{2}/\d{2}/\d{4})\s+([^\n]{0,80}?monetary policy[^\n]{0,60})", txt, re.I)
        report("ecb calendario reuniones", r.ok and bool(hits), f"HTTP {r.status_code}; {len(hits)}: "
               + " | ".join(f"{d} {re.sub(r'\\s+', ' ', t)[:70]}" for d, t in hits[:6]))
    except Exception as exc:
        report("ecb calendario reuniones", False, str(exc)[:200])
    feeds = {
        "fed": "https://www.federalreserve.gov/feeds/press_all.xml",
        "ecb": "https://www.ecb.europa.eu/rss/press.html",
        "bls": "https://www.bls.gov/feed/bls_latest.rss",
        "imf": "https://www.imf.org/en/News/RSS?Language=ENG",
        "sec-rtx": "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=0000101829&type=8-K&dateb=&owner=include&count=10&output=atom",
        "boe": "https://www.bankofengland.co.uk/rss/news",
    }
    sec_ua = {"User-Agent": "market-radar personal research ohlala-radar@example.com"}
    for name, url in feeds.items():
        try:
            r = requests.get(url, headers=sec_ua if name.startswith("sec") else UA, timeout=60)
            items = re.findall(r"<(?:item|entry)\b", r.text)
            titles = re.findall(r"<title[^>]*>(?:<!\[CDATA\[)?([^<\]]{5,120})", r.text)[1:3]
            dates = re.findall(r"<(?:pubDate|updated|dc:date)>([^<]+)<", r.text)[:1]
            report(f"rss {name}", r.ok and bool(items), f"HTTP {r.status_code} {r.headers.get('content-type', '')[:30]}, "
                   f"{len(items)} entradas; {titles} {dates}")
        except Exception as exc:
            report(f"rss {name}", False, str(exc)[:200])


def flows(part: str) -> None:
    """COT de la CFTC, derivados cripto, flujos de ETF y volumen."""
    import json as _json
    if part == "crypto":
        _crypto(_json)
        return
    if part == "etf":
        _etf()
        return
    for ds, code, name in [("72hh-3qpy", "088691", "disagg oro"), ("gpe5-46if", "13874A", "tff sp500"),
                           ("6dca-aqww", "133741", "legacy btc cme")]:
        try:
            r = requests.get(f"https://publicreporting.cftc.gov/resource/{ds}.json", headers=UA, timeout=60, params={
                "$where": f"cftc_contract_market_code='{code}'", "$order": "report_date_as_yyyy_mm_dd DESC", "$limit": 2})
            rows = r.json() if r.ok else []
            keys = sorted(rows[0].keys()) if rows else []
            report(f"cot {name}", bool(rows), f"HTTP {r.status_code}; fecha {rows[0].get('report_date_as_yyyy_mm_dd') if rows else ''}; "
                   f"{rows[0].get('market_and_exchange_names') if rows else r.text[:150]}; campos: {', '.join(k for k in keys if 'pct' not in k and 'conc' not in k and 'traders' not in k)[:600]}")
        except Exception as exc:
            report(f"cot {name}", False, str(exc)[:200])
    _etf()


def _crypto(_json) -> None:
    crypto = {
        "deribit summary": "https://www.deribit.com/api/v2/public/get_book_summary_by_instrument?instrument_name=BTC-PERPETUAL",
        "deribit funding": "https://www.deribit.com/api/v2/public/get_funding_rate_history?instrument_name=BTC-PERPETUAL&start_timestamp=1760000000000&end_timestamp=1760600000000",
        "okx oi": "https://www.okx.com/api/v5/rubik/stat/contracts/open-interest-volume?ccy=BTC&period=1D",
        "okx funding": "https://www.okx.com/api/v5/public/funding-rate-history?instId=BTC-USDT-SWAP&limit=3",
        "okx liq": "https://www.okx.com/api/v5/public/liquidation-orders?instType=SWAP&uly=BTC-USDT&state=filled&limit=3",
        "bybit oi": "https://api.bybit.com/v5/market/open-interest?category=linear&symbol=BTCUSDT&intervalTime=1d&limit=3",
        "binance oi": "https://fapi.binance.com/futures/data/openInterestHist?symbol=BTCUSDT&period=1d&limit=3",
        "hyperliquid": "https://api.hyperliquid.xyz/info",
    }
    for name, url in crypto.items():
        try:
            if name == "hyperliquid":
                r = requests.post(url, json={"type": "metaAndAssetCtxs"}, headers=UA, timeout=30)
                body = r.json() if r.ok else r.text
                ctx = body[1][0] if r.ok else body
                report(name, r.ok, f"HTTP {r.status_code} {_json.dumps(ctx)[:300]}")
                continue
            r = requests.get(url, headers=UA, timeout=30)
            report(f"cripto {name}", r.ok, f"HTTP {r.status_code} {r.text[:350]}")
        except Exception as exc:
            report(f"cripto {name}", False, str(exc)[:200])


def _etf() -> None:
    etf = {
        "spy ssga navhist": "https://www.ssga.com/us/en/intermediary/library-content/products/fund-data/etfs/us/navhist-us-en-spy.xlsx",
        "gld ssga navhist": "https://www.ssga.com/us/en/intermediary/library-content/products/fund-data/etfs/us/navhist-us-en-gld.xlsx",
        "ibit xls": "https://www.ishares.com/us/products/333011/ishares-bitcoin-trust-etf/1467271812596.ajax?fileType=xls&fileName=iShares-Bitcoin-Trust-ETF_fund&dataType=fund",
        "iau xls": "https://www.ishares.com/us/products/239561/ishares-gold-trust-fund/1467271812596.ajax?fileType=xls&fileName=iShares-Gold-Trust_fund&dataType=fund",
    }
    for name, url in etf.items():
        try:
            r = requests.get(url, headers=UA, timeout=60)
            ct = r.headers.get("content-type", "")
            body = r.text[:300].replace("\n", " ") if "text" in ct or "json" in ct or "csv" in ct else f"<{len(r.content)} bytes>"
            if "ishares" in url and r.ok:
                import re as _re
                t = r.content.decode("utf-8", "ignore")
                sheets = _re.findall(r'ss:Name="([^"]+)"', t)
                i = t.find("Shares Outstanding")
                body = f"hojas {sheets}; len {len(t)}; inicio {t[:200]!r}; SO@{i}: {t[max(0, i - 300):i + 600]!r}"
            if url.endswith(".xlsx") and r.ok:
                import io as _io
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(_io.BytesIO(r.content), read_only=True)
                    ws = wb.worksheets[0]
                    rows = [[c for c in row] for row in ws.iter_rows(max_row=6, values_only=True)]
                    body = f"hojas {wb.sheetnames}; {rows}"
                except Exception as e2:
                    body += f" (openpyxl: {e2})"
            report(f"etf {name}", r.ok, f"HTTP {r.status_code} {ct[:40]} {body}")
        except Exception as exc:
            report(f"etf {name}", False, str(exc)[:200])
    key = os.environ.get("TIINGO_API_KEY", "").strip()
    try:
        r = requests.get("https://api.tiingo.com/tiingo/daily/GLD/prices", headers={"Authorization": f"Token {key}"},
                         params={"startDate": START}, timeout=30)
        d = r.json()
        report("tiingo volumen GLD", r.ok, f"último {d[-1]['date'][:10]} volume {d[-1].get('volume')}" if r.ok and d else r.text[:200])
    except Exception as exc:
        report("tiingo volumen", False, str(exc)[:200])


def main() -> int:
    group = sys.argv[1] if len(sys.argv) > 1 else "tiingo"
    if group == "tiingo":
        for pair in ("xauusd", "xagusd"):
            tiingo_fx(pair)
        tiingo_fx_history("xauusd")
    elif group == "imf":
        imf()
    elif group == "events":
        events()
    elif group in ("flows", "crypto", "etf"):
        flows(group)
    elif group == "public":
        euronext("FR0000121329-XPAR")
    return 0


if __name__ == "__main__":
    sys.exit(main())
